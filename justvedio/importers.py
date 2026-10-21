import re
from pathlib import Path

from .models import Candidate, IdentityImport


def load_identities(path: str) -> IdentityImport:
    from openpyxl import load_workbook

    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 身份文件")
    wb = load_workbook(path, read_only=True, data_only=True)
    required = ("所在的区县", "姓名", "身份证号")
    try:
        for sheet in wb:
            rows = sheet.iter_rows()
            header = next(rows, ())
            names = [str(c.value or "").strip() for c in header]
            if not all(n in names for n in required):
                continue
            indices = [names.index(n) for n in required]
            candidates, issues, rejected = [], [], []
            seen = set()
            for row_number, cells in enumerate(rows, 2):
                chosen = [cells[i].value if i < len(cells) else None for i in indices]
                if all(v is None or str(v).strip() == "" for v in chosen):
                    continue
                district, name, raw_id = chosen
                def reject(reason):
                    issues.append(f"第 {row_number} 行：{reason}")
                    values = ["" if v is None else str(v).strip() for v in chosen]
                    rejected.append((Candidate(row_number, *values), reason))
                if any(v is None or not str(v).strip() for v in chosen):
                    reject("区县、姓名、身份证号均为必填")
                    continue
                if not isinstance(raw_id, str):
                    # Excel keeps only 15 significant decimal digits. Never guess lost ID digits.
                    if isinstance(raw_id, bool) or not isinstance(raw_id, (int, float)) or raw_id != int(raw_id) or len(str(int(raw_id))) != 15:
                        reject("身份证号为数值或格式异常，请从原始资料重新录入为文本；无法恢复已丢失的精度")
                        continue
                    raw_id = str(int(raw_id))
                id_number = raw_id.strip().upper()
                if not re.fullmatch(r"(?:[0-9]{15}|[0-9]{17}[0-9X])", id_number):
                    reject("身份证号应为 15 位或 18 位，末位可为 X")
                    continue
                if id_number in seen:
                    reject("身份证号重复，避免同一身份同时登录")
                    continue
                seen.add(id_number)
                candidates.append(Candidate(row_number, str(district).strip(), str(name).strip(), id_number))
            return IdentityImport(candidates, issues, sheet.title, rejected)
        raise ValueError("找不到包含「所在的区县、姓名、身份证号」三列表头的工作表")
    finally:
        wb.close()
