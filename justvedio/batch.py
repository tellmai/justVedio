"""Bounded concurrent scheduling with durable, re-importable results."""
import asyncio
import json
import os
import threading
import uuid
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from pathlib import Path
from .models import Candidate
from .privacy import redact

LABELS = {"pending": "未开始", "running": "正在登录", "watching": "正在观看", "completed": "流程完成", "failed": "失败", "stopped": "已中止"}
ACTIVE_STATUSES = {'running','watching'}
def status_text(record):
    return '登录验证完成' if record.status=='completed' and record.stage=='登录验证完成' else LABELS[record.status]
def now():
    return datetime.now().isoformat(timespec="seconds")


@dataclass(frozen=True)
class PersonResult:
    index: int
    candidate: Candidate
    status: str = "pending"
    reason: str = ""
    started: str = ""
    ended: str = ""
    progress: float = 0
    videos_completed: int = 0
    elapsed_seconds: float = 0
    stage: str = "登录验证"


class StatisticsError(RuntimeError):
    pass


class BatchStore:
    def __init__(self, root, candidates, rejected=()):
        self.flow_kind = 'video'
        self.execution_settings = None
        self.folder = Path(root).resolve() / f"视频批次_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
        pairs = sorted([(c, "") for c in candidates] + list(rejected), key=lambda p: p[0].row)
        self.records = [PersonResult(i, c, "failed" if reason else "pending", reason, ended=now() if reason else "") for i, (c, reason) in enumerate(pairs)]

    def update(self, index, persist=True, **changes):
        self.records[index] = replace(self.records[index], **changes)
        if persist:
            self.persist()
        return self.records[index]

    def persist(self):
        try:
            self.folder.mkdir(parents=True, exist_ok=True)
            target = self.folder / "人员状态.json"
            temporary = target.with_suffix(".json.tmp")
            with temporary.open("w", encoding="utf-8") as stream:
                json.dump({"version": 3, "application": "justVedio", "execution_settings": self.execution_settings, "test_mode":self.flow_kind=='test', "stage": "登录验证" if self.flow_kind=='login' else "登录→视频观看", "updated": now(), "records": [asdict(r) for r in self.records]}, stream, ensure_ascii=False, indent=2)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, target)
        except OSError as error:
            raise StatisticsError("人员状态保存失败，队列已停止，请检查结果目录和磁盘空间") from error
        self.export()

    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        try:
            for completed, filename in ((True, "已完成信息统计.xlsx"), (False, "未完成信息统计.xlsx")):
                wb = Workbook()
                try:
                    ws = wb.active
                    ws.title = "流程统计"
                    ws.append(["所在的区县", "姓名", "身份证号", "源文件行号", "状态", "处理阶段", "已完成视频数", "当前视频进度(%)", "观看用时(秒)", "开始时间", "结束时间", "原因"])
                    for r in self.records:
                        if (r.status == "completed") != completed:
                            continue
                        c = r.candidate
                        values = [c.district, c.name, c.id_number, c.row, status_text(r), r.stage, r.videos_completed, round(r.progress, 1), round(r.elapsed_seconds, 1), r.started, r.ended, r.reason]
                        ws.append(values)
                        for cell, value in zip(ws[ws.max_row], values):
                            if isinstance(value, str):
                                cell.data_type, cell.number_format = "s", "@"
                    for cell in ws[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="2464AC")
                    for col, width in zip("ABCDEFGHIJKL", (18, 18, 26, 12, 18, 22, 18, 22, 20, 24, 24, 65)):
                        ws.column_dimensions[col].width = width
                    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
                    target = self.folder / filename
                    temporary = target.with_suffix(".tmp.xlsx")
                    wb.save(temporary)
                    os.replace(temporary, target)
                finally:
                    wb.close()
        except (OSError, ValueError) as error:
            raise StatisticsError("Excel 导出失败，队列已停止；请关闭占用的统计表后重新导出，最新结果保存在 JSON 中") from error


class BatchControl:
    def __init__(self):
        self.paused = threading.Event()
        self.after_current = threading.Event()
        self.immediate = threading.Event()


class BatchRunner:
    def __init__(self, store, control, changed=lambda _: None):
        self.store, self.control, self.changed = store, control, changed

    def update(self, index, **changes):
        try:
            return self.store.update(index, **changes)
        except StatisticsError:
            if changes.get("status") == "running":
                self.store.records[index] = replace(self.store.records[index], status="failed", reason="统计写入失败，未开始浏览器操作", ended=now())
            raise
        finally:
            self.changed(self.store.records[index])

    async def run(self, perform, concurrency=10):
        if not 1 <= concurrency <= 30:
            raise ValueError("并发数须为 1–30")
        self.store.persist()
        pending = [r.index for r in self.store.records if r.status == "pending"]
        async def worker():
            while pending:
                while self.control.paused.is_set() and not (self.control.immediate.is_set() or self.control.after_current.is_set()):
                    await asyncio.sleep(0.05)
                if self.control.immediate.is_set() or self.control.after_current.is_set():
                    return
                index = pending.pop(0)
                candidate = self.store.records[index].candidate
                self.update(index, status="running", started=now())
                try:
                    outcome = await perform(candidate, index)
                except asyncio.CancelledError:
                    self.update(index, status="stopped", reason="用户停止或批次中断，观看结果未确认", ended=now())
                    raise
                except StatisticsError:
                    raise
                except Exception as error:
                    from playwright.async_api import Error as PlaywrightError
                    message = "页面超时、浏览器已关闭或页面操作失败" if isinstance(error, PlaywrightError) else str(error)
                    self.update(index, status="failed", reason=redact(message, (candidate.name, candidate.id_number)), ended=now())
                else:
                    login_only=self.store.flow_kind=='login'
                    self.update(index, status="completed", stage="登录验证完成" if login_only else "全部完成", reason="登录已验证，停留在登录后主页" if login_only else "指定流程全部完成", progress=100, ended=now(), **(outcome or {}))
        tasks = [asyncio.create_task(worker()) for _ in range(min(concurrency, len(pending))) ]
        try:
            await asyncio.gather(*tasks)
        except BaseException:
            for task in tasks:
                task.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
            # A persistence failure inside perform must not leave phantom active people.
            for r in self.store.records:
                if r.status in ACTIVE_STATUSES:
                    self.store.update(r.index, persist=False, status="stopped", ended=now(), reason="批次中断，结果未确认")
                    self.changed(self.store.records[r.index])
            try:
                self.store.persist()
            except StatisticsError:
                pass
            raise
        return self.counts()

    def counts(self):
        return {s: sum(r.status == s for r in self.store.records) for s in LABELS}
