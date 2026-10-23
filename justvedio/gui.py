import logging
import sys
from datetime import datetime
from pathlib import Path
from PySide6.QtCore import QThread, QTimer, Qt, Signal, QUrl
from PySide6.QtGui import QDesktopServices, QIcon
from PySide6.QtWidgets import (QAbstractItemView, QCheckBox, QComboBox, QFileDialog,
    QFrame, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QPlainTextEdit, QProgressBar, QPushButton, QSizePolicy, QSpinBox,
    QSplitter, QTabWidget, QTableWidgetItem, QToolButton, QVBoxLayout, QWidget,
    QApplication, QScrollArea)
from .controller import BrowserWorker
from .models import DEFAULT_URL, DEFAULT_VIDEO_URL, SiteConfig, VideoOptions
from .importers import load_identities
from .privacy import redact
from .batch import LABELS, ACTIVE_STATUSES, status_text
from .version import APP_NAME, VERSION
from .settings import application_settings
from .expiry import is_expired, EXPIRY_MESSAGE, VALID_THROUGH
from .ui import ElidedLabel, MetricCard, PeopleTable, Switch, icon, stylesheet


class ImportWorker(QThread):
    result = Signal(object)
    failed = Signal(str)
    progress = Signal(str)
    def __init__(self, path, parent=None):
        super().__init__(parent)
        self.path = path
    def run(self):
        try:
            result = load_identities(self.path)
            if self.isInterruptionRequested():
                raise InterruptedError('已停止文件导入')
            self.result.emit(result)
        except Exception as error:
            self.failed.emit(redact(str(error)) if isinstance(error, (ValueError,InterruptedError)) else "文件读取失败，请检查格式和访问权限")


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = application_settings()
        self.candidates, self.rejected, self.batch_records = [], [], []
        self.import_worker = None
        self.busy = self.backend_ready = self._closing = self._allow_close = False
        self.result_folder = ""
        self.batch_test_mode = None
        self._expiry_notified = False
        self.setWindowTitle(f"{APP_NAME} v{VERSION} · 有效期至 {VALID_THROUGH}")
        available = QApplication.primaryScreen().availableGeometry()
        self.setMinimumSize(min(820, available.width()), min(620, available.height()))
        self.resize(min(1180, available.width()), min(900, available.height()))
        resources = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parents[1]))
        self.setWindowIcon(QIcon(str(resources / 'justvedio/assets/app.ico')))
        self._build_ui()
        self.backend = BrowserWorker(self)
        self.backend.ready.connect(self._ready)
        self.backend.progress.connect(self.show_status)
        self.backend.failed.connect(lambda text: self.show_status("失败：" + text))
        self.backend.completed.connect(self.show_status)
        self.backend.batch_started.connect(self._batch_started)
        self.backend.person_updated.connect(self._person_updated)
        self.backend.batch_finished.connect(lambda _: self._render())
        self.backend.job_finished.connect(self._finished)
        self.backend.check_result.connect(lambda label, ok, detail: self.show_status(f"{label}：{'通过' if ok else '未通过'} · {detail}"))
        self.backend.finished.connect(self._backend_finished)
        self.backend.start()
        self._controls()
        self.expiry_timer=QTimer(self)
        self.expiry_timer.timeout.connect(self._check_expiry)
        self.expiry_timer.start(1000)

    def _check_expiry(self):
        if not is_expired():
            return False
        if not self._expiry_notified:
            self._expiry_notified=True
            self.expiry_timer.stop()
            self.close()
            QMessageBox.critical(self,'软件已过期',EXPIRY_MESSAGE)
        return True

    def _build_ui(self):
        root = QWidget()
        root.setObjectName('workspace')
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(22, 18, 22, 18)
        layout.setSpacing(16)

        def label(text, name, wrap=False):
            result = QLabel(text)
            result.setObjectName(name)
            result.setWordWrap(wrap)
            return result

        def field(target, caption, control, hint=''):
            target.addWidget(label(caption, 'fieldLabel'))
            target.addWidget(control)
            if hint:
                target.addWidget(label(hint, 'fieldHint', True))
            target.addSpacing(8)

        def path_row(edit, button):
            container = QWidget()
            row = QHBoxLayout(container)
            row.setContentsMargins(0, 0, 0, 0)
            row.setSpacing(6)
            edit.setMinimumWidth(0)
            row.addWidget(edit, 1)
            row.addWidget(button)
            return container

        def page():
            content = QWidget()
            body = QVBoxLayout(content)
            body.setContentsMargins(0, 16, 8, 6)
            body.setSpacing(8)
            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setFrameShape(QFrame.NoFrame)
            scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
            scroll.setWidget(content)
            return content, body, scroll

        header = QHBoxLayout()
        header.setSpacing(13)
        mark = QLabel()
        mark.setFixedSize(46, 46)
        mark.setPixmap(icon('brand', size=46).pixmap(46, 46))
        header.addWidget(mark)
        brand = QVBoxLayout()
        brand.setSpacing(1)
        brand_line = QHBoxLayout()
        brand_line.setSpacing(10)
        brand_line.addWidget(label(APP_NAME, 'brandTitle'))
        brand_line.addWidget(label(f'v{VERSION}', 'versionBadge'), 0, Qt.AlignVCenter)
        brand_line.addStretch()
        brand.addLayout(brand_line)
        brand.addWidget(label('视频任务工作台  /  登录 · 观看 · 自动补位', 'brandSubtitle'))
        header.addLayout(brand, 1)
        system = QVBoxLayout()
        system.setSpacing(5)
        self.run_badge = label('●  正在初始化', 'runBadge')
        system.addWidget(self.run_badge, 0, Qt.AlignRight)
        system.addWidget(label(f'有效期至 {VALID_THROUGH}', 'expiryCaption'), 0, Qt.AlignRight)
        header.addLayout(system)
        layout.addLayout(header)

        self.workspace_splitter = QSplitter(Qt.Horizontal)
        self.workspace_splitter.setChildrenCollapsible(False)
        self.workspace_splitter.setHandleWidth(16)
        layout.addWidget(self.workspace_splitter, 1)
        config = QFrame()
        config.setObjectName('configPanel')
        config.setMinimumWidth(276)
        config.setMaximumWidth(380)
        config_layout = QVBoxLayout(config)
        config_layout.setContentsMargins(18, 18, 10, 14)
        config_layout.setSpacing(7)
        config_layout.addWidget(label('运行配置', 'panelTitle'))
        config_layout.addWidget(label('确认流程后，即可启动本次任务', 'panelCaption'))
        self.settings_tabs = QTabWidget()
        self.settings_tabs.setObjectName('configTabs')
        self.settings_tabs.setDocumentMode(True)
        config_layout.addWidget(self.settings_tabs, 1)

        execution, execution_layout, self.settings_scroll = page()
        option = QFrame()
        option.setObjectName('modeOption')
        option_layout = QVBoxLayout(option)
        option_layout.setContentsMargins(12, 9, 12, 12)
        option_layout.setSpacing(3)
        self.test_mode = Switch('启动测试')
        self.test_mode.setObjectName('testModeSwitch')
        self.test_mode.setChecked(False)
        self.test_mode.setToolTip('仅切换执行流程：关闭时只验证登录并保留主页；开启后登录成功直接播放视频。视频设置始终可提前配置。')
        self.test_mode.toggled.connect(self._controls)
        option_layout.addWidget(self.test_mode)
        self.mode_description = label('', 'modeDescription', True)
        option_layout.addWidget(self.mode_description)
        execution_layout.addWidget(option)
        execution_layout.addSpacing(10)
        concurrency_row = QHBoxLayout()
        concurrency_row.addWidget(label('并发数量', 'fieldLabel'))
        concurrency_row.addStretch()
        self.concurrency = QSpinBox()
        self.concurrency.setObjectName('flowConcurrency')
        self.concurrency.setRange(1, 30)
        self.concurrency.setValue(int(self.settings.value('concurrency', 10)))
        self.concurrency.setSuffix(' 路')
        self.concurrency.setFixedWidth(88)
        self.concurrency.setToolTip('每个身份从登录到所选流程结束共占一个名额，完成或失败后自动接入下一位。')
        self.concurrency.valueChanged.connect(lambda _: self._render() if hasattr(self, 'tables') else None)
        concurrency_row.addWidget(self.concurrency)
        execution_layout.addLayout(concurrency_row)
        execution_layout.addWidget(label('每个身份占用 1 路，结束后自动补位。', 'fieldHint', True))
        execution_layout.addSpacing(8)
        divider = QFrame()
        divider.setObjectName('divider')
        divider.setFixedHeight(1)
        execution_layout.addWidget(divider)
        execution_layout.addSpacing(6)
        self.background = Switch('后台执行（隐藏浏览器）')
        self.background.setObjectName('backgroundExecution')
        self.background.setChecked(self.settings.value('background', True, type=bool))
        self.background.setToolTip('仅隐藏本批次浏览器，主界面继续显示进度、日志及暂停/停止操作。取消勾选后可观察浏览器；设置在批次启动时生效。')
        self.background.toggled.connect(self._controls)
        execution_layout.addWidget(self.background)
        execution_layout.addWidget(label('主界面保持可见，可随时暂停或停止。', 'fieldHint', True))
        execution_layout.addSpacing(12)
        workflow = QFrame()
        workflow.setObjectName('workflowCard')
        workflow_layout = QVBoxLayout(workflow)
        workflow_layout.setContentsMargins(12, 12, 12, 14)
        workflow_layout.setSpacing(14)
        workflow_layout.addWidget(label('本次执行流程', 'fieldHint'))
        self.workflow_labels = []
        for number, title, note in ((1, '身份登录', '独立会话 · 自动识别验证码'),
                                    (2, '保留登录主页', '仅登录验证，不播放视频'),
                                    (3, '保存结果并自动补位', '结果保存为 Excel 与 JSON')):
            step = QHBoxLayout()
            step.setSpacing(10)
            number_label = label(str(number).zfill(2), 'stepNumber')
            number_label.setAlignment(Qt.AlignCenter)
            number_label.setFixedSize(22, 22)
            step.addWidget(number_label, 0, Qt.AlignTop)
            copy = QVBoxLayout()
            copy.setSpacing(3)
            step_title = label(title, 'stepTitle')
            step_note = label(note, 'stepNote', True)
            copy.addWidget(step_title)
            copy.addWidget(step_note)
            self.workflow_labels.append((step_title, step_note))
            step.addLayout(copy, 1)
            workflow_layout.addLayout(step)
        execution_layout.addWidget(workflow)
        execution_layout.addStretch()
        execution_layout.addWidget(label('运行中配置将锁定，批次结束后恢复编辑。', 'fieldHint', True))
        self.settings_tabs.addTab(self.settings_scroll, '执行')

        self.flow_settings, video_layout, video_scroll = page()
        self.mode = QComboBox()
        self.mode.addItem('直接打开视频网址', 'url')
        self.mode.addItem('点击站内视频入口', 'link')
        self.mode.setCurrentIndex(int(self.settings.value("mode", 0)))
        field(video_layout, '观看入口', self.mode)
        self.video_url = QLineEdit(self.settings.value("video_url", DEFAULT_VIDEO_URL))
        field(video_layout, '测试视频网址', self.video_url)
        self.entry = QLineEdit(self.settings.value("entry", ""))
        self.entry.setPlaceholderText('例如 text=观看视频')
        field(video_layout, '站内入口选择器', self.entry, '站内具体按钮适配可在后续配置。')
        self.muted = QCheckBox("静音播放")
        self.muted.setChecked(self.settings.value("muted", True, type=bool))
        video_layout.addWidget(self.muted)
        video_layout.addSpacing(8)
        self.advanced_toggle = QCheckBox('高级设置：播放器与超时')
        video_layout.addWidget(self.advanced_toggle)
        advanced = QWidget()
        extra = QVBoxLayout(advanced)
        extra.setContentsMargins(0, 10, 0, 0)
        extra.setSpacing(6)
        self.media = QLineEdit(self.settings.value("media", "video, bwp-video"))
        self.play = QLineEdit(self.settings.value("play", ""))
        self.play.setPlaceholderText('留空时点击视频播放')
        self.frame = QLineEdit(self.settings.value("frame", ""))
        self.frame.setPlaceholderText('可选')
        self.complete = QLineEdit(self.settings.value("complete", ""))
        self.complete.setPlaceholderText('可选')
        self.next_video = QLineEdit(self.settings.value("next", ""))
        self.next_video.setPlaceholderText('可选，用于连续视频')
        self.video_timeout = QSpinBox()
        self.video_timeout.setRange(1, 720)
        self.video_timeout.setValue(int(self.settings.value("video_timeout", 120)))
        self.video_timeout.setSuffix(' 分钟')
        self.stall_timeout = QSpinBox()
        self.stall_timeout.setRange(5, 1800)
        self.stall_timeout.setValue(int(self.settings.value("stall_timeout", 120)))
        self.stall_timeout.setSuffix(' 秒')
        for caption, control in (('视频元素', self.media), ('播放按钮', self.play),
                                 ('视频 iframe', self.frame), ('平台完成标志', self.complete),
                                 ('下一视频按钮', self.next_video), ('单视频最长等待', self.video_timeout),
                                 ('停滞超时', self.stall_timeout)):
            field(extra, caption, control)
        video_layout.addWidget(advanced)
        advanced.hide()
        self.advanced_toggle.toggled.connect(advanced.setVisible)
        self.mode.currentIndexChanged.connect(self._controls)
        video_layout.addStretch()
        video_layout.addWidget(label('开启“启动测试”后生效；以真实播放结束为完成依据。', 'fieldHint', True))
        self.settings_tabs.addTab(video_scroll, '视频')

        connection, connection_layout, connection_scroll = page()
        self.url = QLineEdit(self.settings.value('url', DEFAULT_URL))
        field(connection_layout, '登录入口', self.url)
        self.chrome = QLineEdit(self.settings.value('chrome', ''))
        self.chrome.setPlaceholderText('自动查找 Chrome / Edge')
        self.chrome_button = QPushButton('浏览…')
        self.chrome_button.setToolTip('选择浏览器 EXE')
        self.chrome_button.clicked.connect(self._choose_chrome)
        field(connection_layout, '浏览器', path_row(self.chrome, self.chrome_button), '留空自动发现，可使用随附浏览器。')
        base = Path(sys.executable).parent if getattr(sys, 'frozen', False) else Path(__file__).resolve().parents[1]
        self.output = QLineEdit(self.settings.value('output', str(base / 'results')))
        self.output_button = QPushButton('浏览…')
        self.output_button.setToolTip('选择结果目录')
        self.output_button.clicked.connect(self._choose_output)
        field(connection_layout, '结果目录', path_row(self.output, self.output_button), '每个批次独立保存统计文件，可重新导出。')
        connection_layout.addStretch()
        connection_layout.addWidget(label('每个身份使用独立会话，Cookie 与本地存储相互隔离。', 'fieldHint', True))
        self.settings_tabs.addTab(connection_scroll, '连接')
        self.inspect_button = QPushButton('环境检测')
        self.inspect_button.setIcon(icon('monitor'))
        self.inspect_button.clicked.connect(lambda: self._start('inspect'))
        config_layout.addSpacing(5)
        config_layout.addWidget(self.inspect_button)
        self.workspace_splitter.addWidget(config)

        dashboard = QWidget()
        dashboard.setMinimumWidth(450)
        dashboard_layout = QVBoxLayout(dashboard)
        dashboard_layout.setContentsMargins(0, 0, 0, 0)
        dashboard_layout.setSpacing(14)
        metrics = QHBoxLayout()
        metrics.setSpacing(10)
        self.metrics = {}
        for key, title, mark, tone, note in (
            ('total', '已导入', 'people', 'neutral', '等待导入'),
            ('active', '进行中', 'play', 'blue', '登录与观看'),
            ('completed', '已完成', 'check', 'teal', '完成率 0%'),
            ('failed', '需关注', 'alert', 'amber', '失败或中止')):
            card = MetricCard(title, mark, tone, note)
            self.metrics[key] = card
            metrics.addWidget(card, 1)
        dashboard_layout.addLayout(metrics)

        people = QFrame()
        people.setObjectName('peoplePanel')
        people_layout = QVBoxLayout(people)
        people_layout.setContentsMargins(16, 14, 16, 12)
        people_layout.setSpacing(10)
        people_header = QHBoxLayout()
        people_header.setSpacing(8)
        people_heading = QVBoxLayout()
        people_heading.setSpacing(4)
        people_heading.addWidget(label('人员任务', 'panelTitle'))
        self.identity_summary = ElidedLabel('导入身份表，开始新的批次')
        self.identity_summary.setObjectName('identitySummary')
        people_heading.addWidget(self.identity_summary)
        people_header.addLayout(people_heading, 1)
        self.import_button = QPushButton("导入身份 Excel")
        self.import_button.setObjectName('importIdentities')
        self.import_button.setIcon(icon('upload', '#137E74', 18))
        self.import_button.setShortcut('Ctrl+O')
        self.import_button.setToolTip('导入身份 Excel（Ctrl+O）')
        self.import_button.clicked.connect(self._choose_file)
        self.template_button = QPushButton('保存模板')
        self.template_button.setProperty('role', 'quiet')
        self.template_button.setIcon(icon('file', size=18))
        self.template_button.setToolTip('保存空白身份表模板，身份证列已设为文本')
        self.template_button.clicked.connect(self._save_template)
        people_header.addWidget(self.template_button)
        people_header.addWidget(self.import_button)
        people_layout.addLayout(people_header)
        self.tabs = QTabWidget()
        self.tabs.setObjectName('peopleTabs')
        self.tabs.setDocumentMode(True)
        self.tabs.setMinimumHeight(104)
        self.tables = {}
        for key, caption, empty_title, empty_note in (
            ('all', '全部身份', '等待导入身份', '导入 Excel 身份表，即可开始登录或视频流程。'),
            ('active', '进行中', '暂无正在执行的身份', '启动任务后，登录与观看进度会显示在这里。'),
            ('completed', '已完成', '暂无已完成的身份', '完成所选流程的身份会自动归入此处。'),
            ('failed', '失败 / 中止', '暂无失败或中止的身份', '异常记录会集中显示，方便查看原因并重试。')):
            table = PeopleTable(empty_title, empty_note)
            table.setHorizontalHeaderLabels(['行号', '区县', '姓名', '身份证号（脱敏）', '状态', '步骤进度', '视频数', '结果 / 原因'])
            table.setEditTriggers(QAbstractItemView.NoEditTriggers)
            table.setSelectionBehavior(QAbstractItemView.SelectRows)
            table.setSelectionMode(QAbstractItemView.SingleSelection)
            table.setShowGrid(False)
            table.setWordWrap(False)
            table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
            table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
            table.setMouseTracking(True)
            heading = table.horizontalHeader()
            heading.setSectionResizeMode(QHeaderView.Interactive)
            heading.setMinimumSectionSize(38)
            heading.setStretchLastSection(True)
            for column, width in enumerate((44, 62, 86, 136, 100, 98, 56, 150)):
                table.setColumnWidth(column, width)
            table.verticalHeader().setDefaultSectionSize(40)
            table.verticalHeader().hide()
            table.cellDoubleClicked.connect(lambda row, col, t=table: self._focus(t, row))
            self.tables[key] = table
            self.tabs.addTab(table, caption)
        people_layout.addWidget(self.tabs, 1)
        progress_row = QHBoxLayout()
        self.counts = label('等待导入', 'queueCaption')
        progress_row.addWidget(self.counts, 1)
        self.progress_caption = label('已完成 0 / 0', 'progressCaption')
        progress_row.addWidget(self.progress_caption)
        self.batch_progress = QProgressBar()
        self.batch_progress.setTextVisible(False)
        self.batch_progress.setRange(0, 100)
        self.batch_progress.setValue(0)
        self.batch_progress.setFixedWidth(76)
        progress_row.addWidget(self.batch_progress)
        people_layout.addLayout(progress_row)
        result_strip = QFrame()
        result_strip.setObjectName('resultStrip')
        result_row = QHBoxLayout(result_strip)
        result_row.setContentsMargins(10, 6, 4, 6)
        result_row.setSpacing(6)
        folder = QLabel()
        folder.setPixmap(icon('folder', '#8A9BA6', 17).pixmap(17, 17))
        result_row.addWidget(folder)
        self.result_location = ElidedLabel('批次结束后，可打开结果目录查看统计文件')
        self.result_location.setObjectName('resultPath')
        result_row.addWidget(self.result_location, 1)
        self.export_button = QPushButton('重新导出')
        self.export_button.setProperty('role', 'quiet')
        self.export_button.clicked.connect(lambda: self._start('export'))
        self.open_button = QPushButton('打开结果目录')
        self.open_button.setProperty('role', 'quiet')
        self.open_button.clicked.connect(lambda: QDesktopServices.openUrl(QUrl.fromLocalFile(self.result_folder)))
        result_row.addWidget(self.export_button)
        result_row.addWidget(self.open_button)
        people_layout.addWidget(result_strip)
        dashboard_layout.addWidget(people, 1)

        log_panel = QFrame()
        log_panel.setObjectName('logPanel')
        log_layout = QVBoxLayout(log_panel)
        log_layout.setContentsMargins(14, 7, 10, 7)
        log_layout.setSpacing(6)
        log_header = QHBoxLayout()
        log_header.setSpacing(8)
        log_icon = QLabel()
        log_icon.setPixmap(icon('log', '#6E8792', 18).pixmap(18, 18))
        log_header.addWidget(log_icon)
        log_header.addWidget(label('运行日志', 'logTitle'))
        self.status = ElidedLabel('正在初始化浏览器控制模块…')
        self.status.setObjectName('liveStatus')
        log_header.addWidget(self.status, 1)
        self.log_toggle = QToolButton()
        self.log_toggle.setCheckable(True)
        self.log_toggle.setText('展开')
        self.log_toggle.setArrowType(Qt.DownArrow)
        self.log_toggle.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.log_toggle.setShortcut('Ctrl+L')
        self.log_toggle.setToolTip('展开 / 收起运行日志（Ctrl+L）')
        log_header.addWidget(self.log_toggle)
        log_layout.addLayout(log_header)
        self.log = QPlainTextEdit()
        self.log.setReadOnly(True)
        self.log.setMaximumBlockCount(1000)
        self.log.setFixedHeight(92)
        self.log.hide()
        self.log_toggle.toggled.connect(self._toggle_log)
        log_layout.addWidget(self.log)
        dashboard_layout.addWidget(log_panel)
        self.workspace_splitter.addWidget(dashboard)
        self.workspace_splitter.setStretchFactor(0, 0)
        self.workspace_splitter.setStretchFactor(1, 1)
        self.workspace_splitter.setSizes([306, 814])

        command = QFrame()
        command.setObjectName('commandBar')
        command_layout = QHBoxLayout(command)
        command_layout.setContentsMargins(18, 12, 14, 12)
        command_layout.setSpacing(18)
        command_copy = QVBoxLayout()
        command_copy.setSpacing(5)
        self.command_title = label('登录验证模式', 'commandTitle')
        self.command_note = ElidedLabel()
        self.command_note.setObjectName('commandNote')
        command_copy.addWidget(self.command_title)
        command_copy.addWidget(self.command_note)
        command_layout.addLayout(command_copy, 1)
        row = self.action_bar = QHBoxLayout()
        row.setSpacing(8)
        self.pause_button = QPushButton("暂停补位")
        self.pause_button.setIcon(icon('pause', size=16))
        self.stop_button = QPushButton("停止")
        self.stop_button.setObjectName('stopFlow')
        self.stop_button.setIcon(icon('stop', '#AD6262', 16))
        self.start_button = QPushButton('启动登录验证')
        self.start_button.setObjectName('startWatch')
        self.start_button.setIcon(icon('play', '#FFFFFF', 16))
        self.start_button.setMinimumWidth(160)
        self.start_button.clicked.connect(lambda: self._start('watch'))
        self.pause_button.clicked.connect(self._pause)
        self.stop_button.clicked.connect(self._stop)
        for button in (self.pause_button, self.stop_button, self.start_button):
            row.addWidget(button)
        command_layout.addLayout(row)
        layout.addWidget(command)
        for button in self.findChildren(QPushButton) + self.findChildren(QToolButton):
            button.setCursor(Qt.PointingHandCursor)
        for edit in (self.url, self.chrome, self.output, self.video_url, self.entry):
            edit.setToolTip(edit.text())
            edit.textChanged.connect(edit.setToolTip)
        self.setStyleSheet(stylesheet())
        self._render()

    def _toggle_log(self, expanded):
        self.log.setVisible(expanded)
        self.log_toggle.setText('收起' if expanded else '展开')
        self.log_toggle.setArrowType(Qt.UpArrow if expanded else Qt.DownArrow)

    def _video_controls(self):
        self.video_url.setEnabled(not self.busy and self.mode.currentData() == "url")
        self.entry.setEnabled(not self.busy and self.mode.currentData() == "link")

    def _controls(self):
        if not hasattr(self,'start_button'):
            return
        locked = self.busy or self.import_worker is not None or self._closing
        for widget in (self.url, self.chrome, self.chrome_button, self.output, self.output_button, self.mode, self.video_url, self.entry, self.concurrency, self.muted, self.media, self.play, self.frame, self.complete, self.next_video, self.video_timeout, self.stall_timeout, self.import_button, self.template_button):
            widget.setEnabled(not locked)
        self.test_mode.setEnabled(not locked)
        self.background.setEnabled(not locked)
        self.flow_settings.setEnabled(not locked)
        self.start_button.setEnabled(self.backend_ready and not locked and bool(self._remaining() or (not self.batch_records and self.rejected)))
        same_batch=bool(self.batch_records) and self.batch_test_mode==self.test_mode.isChecked()
        if self.busy:
            self.start_button.setText('正在执行…')
        elif same_batch and not self._remaining():
            self.start_button.setText('本批次已完成' if all(r.status == 'completed' for r in self.batch_records) else '无可重试身份')
        else:
            self.start_button.setText('重试未完成身份' if same_batch else '启动测试' if self.test_mode.isChecked() else '启动登录验证')
        self.inspect_button.setEnabled(self.backend_ready and not locked)
        self.pause_button.setEnabled(self.busy and not self._closing)
        self.stop_button.setEnabled(self.busy and not self._closing)
        self.export_button.setEnabled(self.backend_ready and not locked and bool(self.result_folder))
        self.open_button.setEnabled(bool(self.result_folder))
        if not locked:
            self._video_controls()
        self._sync_presentation()

    def _sync_presentation(self):
        testing = self.test_mode.isChecked()
        self.mode_description.setText('开启后，登录成功直接观看视频至结束。' if testing else '当前只验证登录，成功后保留主页。')
        title, note = self.workflow_labels[1]
        title.setText('视频观看' if testing else '保留登录主页')
        note.setText('正常播放至结束，再判定完成' if testing else '仅登录验证，不播放视频')
        mode = '视频测试模式' if testing else '登录验证模式'
        execution = '后台执行' if self.background.isChecked() else '显示浏览器'
        self.command_title.setText(f'{mode} · {execution}')
        self.command_note.setText('身份登录 → 正常播放至结束 → 保存结果 → 自动补位' if testing else '完成登录后保留主页；前台模式可双击人员查看网页。')
        if self._closing:
            text, tone = '正在关闭', 'neutral'
        elif self.import_worker is not None:
            text, tone = '正在导入', 'neutral'
        elif not self.backend_ready:
            text, tone = '正在初始化', 'neutral'
        elif self.busy:
            paused = self.backend.control.paused.is_set()
            text, tone = ('已暂停补位', 'amber') if paused else ('任务执行中', 'teal')
        elif self.batch_records and any(r.status in {'failed', 'stopped'} for r in self.batch_records):
            text, tone = '有未完成身份', 'amber'
        elif self.batch_records and all(r.status == 'completed' for r in self.batch_records):
            text, tone = '批次已完成', 'teal'
        else:
            text, tone = '系统就绪', 'teal'
        self.run_badge.setText('●  ' + text)
        if self.run_badge.property('tone') != tone:
            self.run_badge.setProperty('tone', tone)
            self.run_badge.style().unpolish(self.run_badge)
            self.run_badge.style().polish(self.run_badge)

    def _remaining(self):
        if not self.batch_records or self.batch_test_mode!=self.test_mode.isChecked():
            return self.candidates
        valid = {c.row for c in self.candidates}
        return [r.candidate for r in self.batch_records if r.status != "completed" and r.candidate.row in valid]

    def _ready(self):
        self.backend_ready = True
        self.show_status("就绪：关闭“启动测试”时仅验证登录并保留主页；开启后登录成功直接观看视频")
        self._controls()

    def _backend_finished(self):
        self.backend_ready = False
        self._controls()

    def show_status(self, text):
        text = redact(text)
        self.status.setText(text)
        self.log.appendPlainText(f"{datetime.now():%H:%M:%S}  {text}")
        logging.getLogger("justvedio").info(text)

    def _choose_chrome(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择浏览器", "", "程序 (*.exe)")
        if path:
            self.chrome.setText(path)
    def _choose_output(self):
        path = QFileDialog.getExistingDirectory(self, "选择结果目录", self.output.text())
        if path:
            self.output.setText(path)
    def _choose_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "导入身份", "", "Excel (*.xlsx)")
        if path:
            self.import_file(path)
    def import_file(self, path):
        self.import_worker = ImportWorker(path, self)
        self.import_worker.result.connect(self.identities_loaded)
        self.import_worker.failed.connect(self.show_status)
        self.import_worker.finished.connect(self._import_finished)
        self.import_worker.start()
        self._controls()
    def _import_finished(self):
        self.import_worker.deleteLater()
        self.import_worker = None
        self._controls()
    def identities_loaded(self, result):
        self.candidates, self.rejected, self.batch_records = result.candidates, result.rejected, []
        self.batch_test_mode = None
        self.identity_summary.setText(f"工作表 {result.sheet} · 可用 {len(result.candidates)} 人 · 拒绝 {len(result.rejected)} 行")
        for issue in result.issues:
            self.show_status(issue)
        self._render()
        self._controls()
    def _save_template(self):
        path, _ = QFileDialog.getSaveFileName(self, "保存空白身份表模板", "身份模板.xlsx", "Excel (*.xlsx)")
        if path:
            try:
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.append(["所在的区县", "姓名", "身份证号"])
                ws.column_dimensions['A'].width, ws.column_dimensions['B'].width, ws.column_dimensions['C'].width = 20, 18, 28
                for row in ws.iter_rows(min_row=2, max_row=1001, min_col=3, max_col=3):
                    row[0].number_format = '@'
                wb.save(path)
                wb.close()
                self.show_status("身份表模板已保存，身份证列已设为文本")
            except OSError:
                self.show_status("模板保存失败，请检查目标文件是否被占用")

    def _config(self):
        config = SiteConfig(self.url.text().strip(), self.chrome.text().strip())
        options = VideoOptions(concurrency=self.concurrency.value(), background=self.background.isChecked(), video_url=self.video_url.text().strip() if self.mode.currentData() == 'url' else '', entry_selector=self.entry.text().strip() if self.mode.currentData() == 'link' else '', video_selector=self.media.text().strip(), play_selector=self.play.text().strip(), frame_selector=self.frame.text().strip(), complete_selector=self.complete.text().strip(), next_selector=self.next_video.text().strip(), muted=self.muted.isChecked(), video_timeout_seconds=self.video_timeout.value()*60, stall_timeout_seconds=self.stall_timeout.value())
        return config, options
    def _start(self, kind):
        if self._check_expiry():
            return
        try:
            config, options = self._config()
            config.validate()
            options.validate_execution()
            if kind == 'watch' and self.test_mode.isChecked():
                options.validate(config)
                if self.mode.currentData() == 'link' and not options.entry_selector:
                    raise ValueError("请填写原网站内的视频入口选择器")
            if not self.output.text().strip():
                raise ValueError("请填写结果目录")
        except ValueError as error:
            self.show_status(str(error))
            return
        for key, value in {'url': config.url, 'chrome': config.chrome_path, 'output': self.output.text(), 'mode': self.mode.currentIndex(), 'video_url': self.video_url.text(), 'entry': self.entry.text(), 'concurrency': options.concurrency, 'background': options.background, 'muted': options.muted, 'media': options.video_selector, 'play': options.play_selector, 'frame': options.frame_selector, 'complete': options.complete_selector, 'next': options.next_selector, 'video_timeout': self.video_timeout.value(), 'stall_timeout': self.stall_timeout.value()}.items():
            self.settings.setValue(key, value)
        candidates = self._remaining() if kind == 'watch' else self.candidates
        rejected = self.rejected if kind=='watch' and (not self.batch_records or self.batch_test_mode!=self.test_mode.isChecked()) else ()
        self.busy = True
        self.pause_button.setText("暂停补位")
        self.backend.start_job(kind, config, candidates, self.output.text().strip(), rejected, options, test_mode=self.test_mode.isChecked())
        self._controls()
    def _pause(self):
        if self.backend.control.paused.is_set():
            self.backend.resume_queue()
            self.pause_button.setText("暂停补位")
            self.show_status("已恢复自动补位")
        else:
            self.backend.pause_queue()
            self.pause_button.setText("恢复补位")
            self.show_status("已暂停新身份补位；当前身份继续登录或观看")
        self._controls()
    def _stop(self):
        self.backend.pause_queue()
        box = QMessageBox(self)
        box.setWindowTitle("停止流程")
        box.setText("可等待当前所有身份完成登录及视频，或立即关闭当前会话。")
        graceful = box.addButton("当前全部完成后停止", QMessageBox.AcceptRole)
        immediate = box.addButton("立即停止", QMessageBox.DestructiveRole)
        box.addButton("取消", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() == graceful:
            self.backend.stop_after_current()
            self.show_status("停止接入后续身份，等待当前所有身份完成")
        elif box.clickedButton() == immediate:
            self.backend.stop_job()
        elif self.pause_button.text() != "恢复补位":
            self.backend.resume_queue()
    def _finished(self):
        self.busy = False
        self.pause_button.setText("暂停补位")
        self._render()
        self._controls()
    def _batch_started(self, folder, records):
        self.result_folder, self.batch_records = folder, list(records)
        self.batch_test_mode=self.test_mode.isChecked()
        self.result_location.setText(folder)
        self._render()
    def _person_updated(self, record):
        if record.index < len(self.batch_records):
            previous = self.batch_records[record.index]
            self.batch_records[record.index] = record
            if previous.status == record.status:
                for key, table in self.tables.items():
                    row = getattr(self, '_table_rows', {}).get(key, {}).get(record.index)
                    if row is not None:
                        for col, value in ((5, f'{record.progress:.1f}%'), (6, str(record.videos_completed)), (7, record.reason)):
                            table.setItem(row, col, self._table_item(value, record, col))
                return
        self._render()

    @staticmethod
    def _table_item(value, record, column):
        item = QTableWidgetItem(value)
        item.setData(Qt.UserRole, record.index)
        item.setData(Qt.UserRole + 1, record.status)
        item.setToolTip(value)
        if column in (0, 6):
            item.setTextAlignment(Qt.AlignCenter)
        return item

    def _render(self):
        from .batch import PersonResult
        records = self.batch_records or [PersonResult(i, c) for i, c in enumerate(self.candidates)]
        counts = {s: sum(r.status == s for r in records) for s in LABELS}
        total = len(records)
        completed = counts['completed']
        percent = round(completed / total * 100) if total else 0
        values = {'total': total, 'active': counts['running'] + counts['watching'],
                  'completed': completed, 'failed': counts['failed'] + counts['stopped']}
        notes = {'total': f"待接入 {counts['pending']} 人", 'active': f"登录 {counts['running']} · 观看 {counts['watching']}",
                 'completed': f'完成率 {percent}%', 'failed': f"失败 {counts['failed']} · 中止 {counts['stopped']}"}
        for key, card in self.metrics.items():
            card.value.setText(str(values[key]))
            card.note.setText(notes[key])
        self.counts.setText(f"待接入 {counts['pending']} 人 · 并发上限 {self.concurrency.value()}" if records else '导入后可查看各身份的实时进度')
        self.progress_caption.setText(f'已完成 {completed} / {total}')
        self.batch_progress.setValue(percent)
        self.batch_progress.setToolTip(f'已完成 {completed} 人，共 {total} 人（{percent}%）')
        self._table_rows = {}
        captions = {'all': '全部身份', 'active': '进行中', 'completed': '已完成', 'failed': '失败 / 中止'}
        for tab_index, (key, table) in enumerate(self.tables.items()):
            subset = [r for r in records if key == 'all' or (key == 'active' and r.status in ACTIVE_STATUSES) or (key == 'completed' and r.status == 'completed') or (key == 'failed' and r.status in {'failed','stopped'})]
            self.tabs.setTabText(tab_index, f'{captions[key]}  {len(subset)}')
            selected = table.item(table.currentRow(), 0)
            selected_index = selected.data(Qt.UserRole) if selected else None
            vertical, horizontal = table.verticalScrollBar().value(), table.horizontalScrollBar().value()
            table.setUpdatesEnabled(False)
            try:
                table.setRowCount(len(subset))
                self._table_rows[key] = {r.index: row for row, r in enumerate(subset)}
                for row, r in enumerate(subset):
                    c = r.candidate
                    for col, value in enumerate((str(c.row), c.district, c.name, c.masked_id, status_text(r), f'{r.progress:.1f}%', str(r.videos_completed), r.reason)):
                        table.setItem(row, col, self._table_item(value, r, col))
                if selected_index in self._table_rows[key]:
                    table.selectRow(self._table_rows[key][selected_index])
                table.verticalScrollBar().setValue(vertical)
                table.horizontalScrollBar().setValue(horizontal)
            finally:
                table.setUpdatesEnabled(True)
        if hasattr(self, 'command_title'):
            self._sync_presentation()
    def _focus(self, table, row):
        item = table.item(row, 0)
        if item and self.backend_ready:
            self.backend.focus(item.data(Qt.UserRole))
    def closeEvent(self, event):
        if self._allow_close:
            event.accept()
            return
        event.ignore()
        if self._closing:
            return
        self._closing = True
        self.expiry_timer.stop()
        if self.import_worker:
            self.import_worker.requestInterruption()
        self.backend.stop_job()
        self.backend.shutdown()
        self._controls()
        self.show_status("正在关闭本软件创建的浏览器并保存状态…")
        self.close_timer = QTimer(self)
        self.close_timer.timeout.connect(self._finish_close)
        self.close_timer.start(100)
    def _finish_close(self):
        if not self.backend.isRunning() and self.import_worker is None:
            self.close_timer.stop()
            self._allow_close = True
            self.close()
            QApplication.instance().quit()
