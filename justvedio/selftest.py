"""Explicit source/portable checks; synthetic identities stay on loopback."""
import asyncio
import json
import sys
import time
from dataclasses import replace
from datetime import timedelta
from pathlib import Path


def wait_for(app, condition, seconds=90):
    deadline = time.monotonic() + seconds
    while not condition() and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(.01)
    if not condition():
        raise RuntimeError('离线流程等待超时')


def check_gui(window, app, output):
    from PySide6.QtWidgets import QPushButton, QLabel
    from .models import Candidate, IdentityImport
    from .version import APP_NAME
    from .expiry import is_expired, EXPIRES_AT
    people = [Candidate(i+2, '黔江区', f'离线人员{i}', f'11010119900101{i:04}') for i in range(12)]
    window.identities_loaded(IdentityImport(people, [], '合成身份'))
    window.concurrency.setValue(10)
    window.background.setChecked(True)
    window.test_mode.setChecked(False)
    login = window.start_button.isEnabled() and window.start_button.text() == '启动登录验证'
    app.processEvents()
    window.grab().save(str(output/'登录模式预览.png'))
    window.test_mode.setChecked(True)
    app.processEvents()
    test = window.start_button.isEnabled() and window.start_button.text() == '启动测试'
    window.grab().save(str(output/'界面预览.png'))
    texts = [w.text() for kind in (QPushButton, QLabel) for w in window.findChildren(kind)]
    return dict(brand_matches=window.windowTitle().startswith(APP_NAME), login_only_available=login,
                test_without_bank=test, background_checkbox=window.background.text() == '后台执行（隐藏浏览器）',
                no_tray=not hasattr(window, 'tray') and not hasattr(window, 'background_button'),
                no_answering_ui=not any(any(term in text for term in ('答题', '题库', '评分')) for text in texts),
                expiry_enforced=not is_expired(EXPIRES_AT-timedelta(microseconds=1)) and is_expired(EXPIRES_AT))


def run_runtime_checks(output, desktop=False, window=None):
    from openpyxl import Workbook, load_workbook
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QTimer
    from .models import Candidate, VideoOptions, SiteConfig
    from .importers import load_identities
    from .ocr import LocalOCR
    from .fixtures import SyntheticSite
    from .controller import BrowserWorker
    from .browsers import launch_browser
    from .browser import SiteSession
    app = QApplication.instance() or QApplication([])
    people = [Candidate(i+2, '黔江区', f'离线人员{i}', f'11010119900101{i:04}') for i in range(12)]
    wb = Workbook()
    wb.active.append(['所在的区县', '姓名', '身份证号'])
    for person in people:
        wb.active.append([person.district, person.name, person.id_number])
    identity_file = output/'selftest-identities.xlsx'
    wb.save(identity_file)
    wb.close()
    if load_identities(str(identity_file)).candidates != people:
        raise RuntimeError('身份导入自检失败')
    ocr = LocalOCR()
    try:
        asyncio.run(ocr.check())
    finally:
        ocr.close()
    base = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parents[1]))
    bundled = base/'chrome/chrome-win64/chrome.exe'
    if not bundled.is_file():
        bundled = base/'tools/chrome/chrome-win64/chrome.exe'
    media = base/'justvedio/assets/selftest.webm'
    worker = window.backend if window is not None else BrowserWorker()
    ready, finished, errors, transitions, cookies, peaks = [], [], [], {}, [], []
    worker.ready.connect(lambda: ready.append(True))
    worker.job_finished.connect(lambda: finished.append(True))
    worker.failed.connect(errors.append)
    def changed(record):
        chain = transitions.setdefault(record.index, [])
        if not chain or chain[-1] != record.status:
            chain.append(record.status)
    worker.person_updated.connect(changed)
    async def recognize(_):
        return 'TEST'
    original_recognize = worker.ocr.recognize
    worker.ocr.recognize = recognize
    from . import controller
    original_launch, original_login = controller.launch_browser, SiteSession.login
    launch_modes, browser_pids = [], set()
    async def launch(p, explicit='', headless=False, progress=lambda _: None):
        launch_modes.append(headless)
        browser, choice = await launch_browser(p, explicit, headless, progress)
        if desktop:
            cdp = await browser.new_browser_cdp_session()
            try:
                info = await cdp.send('SystemInfo.getProcessInfo')
                browser_pids.add(next(int(item['id']) for item in info['processInfo'] if item['type'] == 'browser'))
            finally:
                await cdp.detach()
        return browser, choice
    async def login(session, *args, **kwargs):
        result = await original_login(session, *args, **kwargs)
        values = await session.context.cookies()
        cookies.append(next(c['value'] for c in values if c['name'] == 'person'))
        peaks.append(len(worker.browser.contexts))
        return result
    controller.launch_browser, SiteSession.login = launch, login
    owns_worker = window is None
    samples, violations, visible_browser_samples = [], [], []
    timer = None
    background_phase = True
    try:
        if owns_worker:
            worker.start()
            wait_for(app, lambda: ready)
        if desktop:
            if window is None:
                raise RuntimeError('桌面验收需要可见的主界面')
            import ctypes
            from ctypes import wintypes
            native = ctypes.WinDLL('user32', use_last_error=True)
            native.GetForegroundWindow.restype = wintypes.HWND
            native.GetWindowThreadProcessId.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.DWORD)]
            native.IsWindowVisible.argtypes = [wintypes.HWND]
            callback_type = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
            native.EnumWindows.argtypes = [callback_type, wintypes.LPARAM]
            def visible_windows():
                handles = []
                @callback_type
                def collect(handle, _):
                    owner = wintypes.DWORD()
                    native.GetWindowThreadProcessId(handle, ctypes.byref(owner))
                    if owner.value in browser_pids and native.IsWindowVisible(handle):
                        handles.append(int(handle))
                    return True
                native.EnumWindows(collect, 0)
                return handles
            def sample():
                if not worker._entered:
                    return
                if background_phase:
                    owner = wintypes.DWORD()
                    native.GetWindowThreadProcessId(native.GetForegroundWindow(), ctypes.byref(owner))
                    samples.append(owner.value)
                    if owner.value in browser_pids or visible_windows():
                        violations.append('后台浏览器显示窗口或抢占前台')
                    if not window.isVisible() or not window.pause_button.isEnabled() or not window.stop_button.isEnabled():
                        violations.append('主界面或运行控制不可用')
                    if window.background.isEnabled() or window.concurrency.isEnabled():
                        violations.append('批次运行配置未锁定')
                elif visible_windows():
                    visible_browser_samples.append(True)
            timer = QTimer()
            timer.timeout.connect(sample)
            timer.start(20)
        with SyntheticSite(media) as site:
            config = SiteConfig(site.origin+'/exam/', str(bundled) if bundled.is_file() else '', 20000, True)
            options = VideoOptions(concurrency=10, background=True, video_url=site.origin+'/watch', play_selector='#play',
                                   complete_selector='#completed', video_timeout_seconds=45, stall_timeout_seconds=20)
            if window is not None:
                window.background.setChecked(True)
                window.busy = True
                window._controls()
            worker.start_job('watch', config, people, str(output/'statistics'), options=options, test_mode=True)
            if desktop:
                wait_for(app, lambda: worker.last_store and any(r.status == 'watching' for r in worker.last_store.records))
                if not window.isVisible():
                    raise RuntimeError('后台执行时主界面被隐藏')
                window._pause()
                if not worker.control.paused.is_set():
                    raise RuntimeError('后台执行时无法暂停补位')
                window._pause()
                window.grab().save(str(output/'后台执行主界面.png'))
            wait_for(app, lambda: finished, 120)
            if errors:
                raise RuntimeError('; '.join(errors))
            records = list(worker.last_store.records)
            if not all(r.status == 'completed' and r.videos_completed == 1 for r in records):
                raise RuntimeError('并发登录及视频未全部完成：'+repr([(r.status,r.reason) for r in records]))
            if set(cookies) != {p.id_number for p in people} or max(peaks) > 10:
                raise RuntimeError('身份隔离或并发上限校验失败')
            if any(chain != ['running', 'watching', 'completed'] for chain in transitions.values()):
                raise RuntimeError('流程步骤不符合登录后直接观看视频')
            if len(site.state['watch_visits']) != 12 or not all(n == 1 for _, n in site.state['watch_visits']):
                raise RuntimeError('视频入口未直接衔接同一身份登录主页')
            state = json.loads((worker.last_store.folder/'人员状态.json').read_text(encoding='utf-8'))
            if state['execution_settings'] != {'concurrency': 10, 'background': True} or launch_modes != [True]:
                raise RuntimeError('后台执行配置没有传递到浏览器启动')
            wb = load_workbook(worker.last_store.folder/'已完成信息统计.xlsx', read_only=True)
            try:
                exported = wb.active.max_row == 13 and wb.active.max_column == 12
            finally:
                wb.close()
            foreground_mode = False
            if desktop:
                if violations or not samples or not window.isVisible():
                    raise RuntimeError('后台模式验收失败：'+repr(violations[:3]))
                if not window.background.isEnabled() or not window.concurrency.isEnabled():
                    raise RuntimeError('批次结束后配置未恢复编辑')
                window.grab().save(str(output/'后台完成主界面.png'))
                background_phase = False
                window.background.setChecked(False)
                window.busy = True
                window._controls()
                worker.start_job('watch', config, people[:1], str(output/'foreground-statistics'),
                                 options=replace(options, background=False, concurrency=1), test_mode=True)
                wait_for(app, lambda: len(finished) == 2, 60)
                foreground_mode = (not errors and launch_modes == [True, False] and bool(visible_browser_samples)
                                   and worker.last_store.records[0].status == 'completed' and window.isVisible())
                if not foreground_mode:
                    raise RuntimeError('取消后台执行后的可见浏览器流程失败')
                window.background.setChecked(True)
            return dict(identity_import=True, statistics_export=exported, ocr_model=True,
                        browser_launch=True, bundled_browser=bundled.is_file(), video_playback=True,
                        twelve_full_flows=True, identity_isolation=True, direct_video_transition=True,
                        max_contexts=max(peaks), completed=len(records), background_execution=True,
                        headless_browser=True, desktop_background=desktop, foreground_mode=foreground_mode,
                        main_window_visible=desktop and window.isVisible(), execution_controls_locked=desktop and not violations,
                        power_request_released=not worker.power_guard.active,
                        background_preserved_foreground=desktop and not violations and bool(samples),
                        foreground_samples=len(samples),
                        no_answering_fields=all(not any(k.startswith('exam_') for k in r) for r in state['records']))
    finally:
        if timer:
            timer.stop()
        controller.launch_browser, SiteSession.login = original_launch, original_login
        worker.ocr.recognize = original_recognize
        if owns_worker:
            worker.shutdown()
            wait_for(app, lambda: not worker.isRunning())
