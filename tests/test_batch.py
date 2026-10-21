import asyncio
import json
import pytest
from openpyxl import load_workbook
from justvedio.models import Candidate
from justvedio.batch import BatchStore, BatchControl, BatchRunner, StatisticsError

def people(n=13):
    return [Candidate(i+2, '黔江区', f'演示{i}', f'11010119900101{i:04}') for i in range(n)]

def test_default_ten_concurrent_and_immediate_refill(tmp_path):
    async def scenario():
        store = BatchStore(tmp_path, people())
        runner = BatchRunner(store, BatchControl())
        entered, release = [], [asyncio.Event() for _ in range(13)]
        peak = 0
        active = 0
        async def perform(c, index):
            nonlocal active, peak
            entered.append(index)
            active += 1
            peak = max(peak, active)
            await release[index].wait()
            active -= 1
            return {'videos_completed': 1, 'elapsed_seconds': 2}
        task = asyncio.create_task(runner.run(perform))
        for _ in range(100):
            if len(entered) == 10:
                break
            await asyncio.sleep(.01)
        assert entered == list(range(10))
        release[4].set()
        for _ in range(100):
            if len(entered) == 11:
                break
            await asyncio.sleep(.01)
        assert entered[-1] == 10 and peak == 10
        assert store.records[4].status == 'completed'
        for gate in release:
            gate.set()
        result = await task
        assert result['completed'] == 13 and peak == 10
        assert all(r.videos_completed == 1 for r in store.records)
    asyncio.run(scenario())

@pytest.mark.parametrize('immediate', [False, True])
def test_stop_preserves_all_active_and_pending(tmp_path, immediate):
    async def scenario():
        store, control = BatchStore(tmp_path, people()), BatchControl()
        runner = BatchRunner(store, control)
        entered, gate = [], asyncio.Event()
        async def perform(c, index):
            entered.append(index)
            await gate.wait()
        task = asyncio.create_task(runner.run(perform))
        while len(entered) < 10:
            await asyncio.sleep(.01)
        if immediate:
            control.immediate.set()
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task
        else:
            control.after_current.set()
            gate.set()
            await task
        assert len(entered) == 10
        assert [r.status for r in store.records] == [('stopped' if immediate else 'completed')]*10 + ['pending']*3
        saved = json.loads((store.folder/'人员状态.json').read_text(encoding='utf-8'))
        assert not any(r['status'] in {'running','watching'} for r in saved['records'])
    asyncio.run(scenario())

def test_one_failure_does_not_block_next_and_identity_redacted(tmp_path):
    store = BatchStore(tmp_path, people(4))
    async def perform(c, i):
        if i == 0:
            raise ValueError(c.name+' '+c.id_number+' 登录失败')
    result = asyncio.run(BatchRunner(store, BatchControl()).run(perform, 2))
    assert result['failed'] == 1 and result['completed'] == 3
    assert store.records[0].candidate.id_number not in store.records[0].reason
    assert store.records[0].candidate.name not in store.records[0].reason
    wb = load_workbook(store.folder/'未完成信息统计.xlsx')
    assert wb.active['C2'].data_type == 's'
    assert wb.active['C2'].value == store.records[0].candidate.id_number
    wb.close()

def test_pause_blocks_refill_but_allows_active_finish(tmp_path):
    async def scenario():
        store, control = BatchStore(tmp_path, people(3)), BatchControl()
        entered, gate = [], asyncio.Event()
        async def perform(c, i):
            entered.append(i)
            if i == 0:
                await gate.wait()
        task = asyncio.create_task(BatchRunner(store, control).run(perform, 1))
        while not entered:
            await asyncio.sleep(.01)
        control.paused.set()
        gate.set()
        await asyncio.sleep(.15)
        assert entered == [0] and not task.done()
        control.paused.clear()
        await task
        assert entered == [0,1,2]
    asyncio.run(scenario())

def test_statistics_failure_stops_before_browser_and_can_reexport(tmp_path, monkeypatch):
    store = BatchStore(tmp_path, people(2))
    original = store.export
    def blocked():
        if any(r.status == 'running' for r in store.records):
            raise StatisticsError('文件占用')
        original()
    monkeypatch.setattr(store, 'export', blocked)
    async def unexpected(*_):
        raise AssertionError('must not start')
    with pytest.raises(StatisticsError):
        asyncio.run(BatchRunner(store, BatchControl()).run(unexpected, 1))
    assert all(r.status != 'running' for r in store.records)
    assert store.records[1].status == 'pending'
    monkeypatch.setattr(store, 'export', original)
    store.persist()

def test_export_formula_like_names_remain_text(tmp_path):
    store = BatchStore(tmp_path, [Candidate(2,'黔江区','=1+1','110101199001010010')])
    store.persist()
    wb = load_workbook(store.folder/'未完成信息统计.xlsx')
    assert wb.active['B2'].value == '=1+1' and wb.active['B2'].data_type == 's'
    wb.close()


def test_export_video_only_preserves_values(tmp_path):
    store=BatchStore(tmp_path,people(1))
    store.update(0,status='completed',stage='全部完成',videos_completed=2,progress=100,elapsed_seconds=38,reason='指定流程全部完成')
    wb=load_workbook(store.folder/'已完成信息统计.xlsx')
    try:
        rows=list(wb.active.values)
        headers=rows[0]
        assert len(headers)==12 and headers[:3]==('所在的区县','姓名','身份证号')
        assert not any('答题' in h or '得分' in h for h in headers)
        record=dict(zip(headers,rows[1]))
        assert (record['已完成视频数'],record['观看用时(秒)'])==(2,38)
    finally:
        wb.close()
